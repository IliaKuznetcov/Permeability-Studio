"""Paired-pulse sessions and trace/state exports; no UI dependency."""

from __future__ import annotations
from dataclasses import asdict
from math import isclose
import json
from pathlib import Path
import tempfile
import os
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .calibration_models import CalibrationProject, ROIStatistics
from .calibration_io import (
    save_session as save_recordings,
    load_session as load_recordings,
    _append,
)
from .pulse_capture_models import PulseProject
from .pulse_capture import capture_pulse, pulse_review_notes


def protect(project, path):
    path = Path(path)
    if any(
        path.resolve() == Path(s.source_path).resolve()
        for s in project.snapshots.values()
        if s.source_path
    ):
        raise ValueError("Choose a new file; source recordings cannot be overwritten.")
    return path


def save_pulse_session(project, path):
    """Write an editable session, including snapshots and pulse selections."""
    path = protect(project, path)
    # Reuse the validated recording serializer; no calibration run/schema changes are required.
    with tempfile.TemporaryDirectory() as directory:
        intermediate = Path(directory) / "recordings.json"
        save_recordings(CalibrationProject(snapshots=project.snapshots), intermediate)
        recordings = json.loads(intermediate.read_text(encoding="utf-8"))
    from . import __version__
    payload = {
        "format": "paired-pulse-session",
        "schema_version": 1,
        "application_version": __version__,
        "recordings": recordings,
        "pulses": [asdict(p) for p in project.pulses],
        "pressure_basis": project.pressure_basis,
        "barometric_kpa": project.barometric_kpa,
        "calibrated_volumes_cm3": project.calibrated_volumes_cm3,
        "calibration_provenance": project.calibration_provenance,
        "analysis_settings": project.analysis_settings,
        "analysis_results": project.analysis_results,
        "sample_results": project.sample_results,
    }
    descriptor, name = tempfile.mkstemp(
        dir=path.parent, prefix=".pulse-", suffix=".tmp"
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            json.dump(payload, stream, ensure_ascii=False, allow_nan=False)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def load_pulse_session(path):
    """Restore recordings and rebuild statistics from the saved selection bounds."""
    payload = json.loads(Path(path).read_text(encoding="utf-8"))
    if (
        payload.get("format") != "paired-pulse-session"
        or payload.get("schema_version") != 1
    ):
        raise ValueError(
            "This is not a supported paired-pulse session. Use Volume calculation to open calibration sessions."
        )
    with tempfile.TemporaryDirectory() as directory:
        intermediate = Path(directory) / "recordings.json"
        intermediate.write_text(
            json.dumps(payload["recordings"], allow_nan=False), encoding="utf-8"
        )
        snapshots = load_recordings(intermediate).snapshots
    basis, bar = payload.get("pressure_basis", "Gauge"), payload.get("barometric_kpa")
    if basis not in ("Gauge", "Absolute") or (
        bar is not None and (not np.isfinite(bar) or bar <= 0)
    ):
        raise ValueError("Saved pressure basis/barometric pressure is invalid.")
    volumes = payload.get("calibrated_volumes_cm3", {})
    if volumes and (
        set(volumes) != {"VU", "VD", "DU", "DD"}
        or any(not np.isfinite(v) or v <= 0 for v in volumes.values())
    ):
        raise ValueError("Saved calibrated volumes are invalid.")
    project = PulseProject(
        snapshots=snapshots,
        pressure_basis=basis,
        barometric_kpa=bar,
        calibrated_volumes_cm3=volumes,
        calibration_provenance=payload.get("calibration_provenance", {}),
        analysis_settings=payload.get("analysis_settings", {}),
        analysis_results=payload.get("analysis_results", {}),
        sample_results=payload.get("sample_results", {}),
    )
    for entry in payload["pulses"]:
        snapshot = snapshots[entry["dataset_id"]]
        kwargs = {
            k: entry[k]
            for k in (
                "uid",
                "label",
                "valve_time_s",
                "initial_pressure_bounds",
                "final_pressure_bounds",
                "initial_temperature_bounds",
                "final_temperature_bounds",
                "pre_context_s",
                "sample_id",
                "notes",
            )
        }
        pulse = capture_pulse(
            snapshot, fit_bounds=(entry["fit_start_s"], entry["fit_end_s"]), **kwargs
        )
        # NumPy reductions can differ by a few floating-point bits between
        # machines. Verify every statistic against the retained raw samples,
        # then keep its saved representation only when it agrees to roundoff.
        # This preserves existing fit fingerprints across platforms without
        # trusting edited or obsolete cached state statistics.
        for state in ("initial", "final"):
            for key, computed in getattr(pulse, state).items():
                cached = entry.get(state, {}).get(key, {})
                if _statistics_agree(computed, cached):
                    getattr(pulse, state)[key] = ROIStatistics(**cached)
        if any(p.uid == pulse.uid or p.label == pulse.label for p in project.pulses):
            raise ValueError("Duplicate pulse ID in saved session.")
        project.pulses.append(pulse)
    return project


def _statistics_agree(computed, cached):
    current = asdict(computed)
    if current.keys() != cached.keys():
        return False
    reductions = {"mean", "std", "minimum", "maximum", "slope_per_min"}
    for key, value in current.items():
        old = cached[key]
        if key in reductions:
            if not isinstance(old, (int, float)) or not isclose(
                    value, old, rel_tol=1e-12, abs_tol=1e-14):
                return False
        elif value != old:
            return False
    return True


def _grouped(project, sheet, mode):
    """Write raw samples as adjacent column groups, one group per pulse."""
    column = 1
    for pulse in project.pulses:
        snapshot = project.snapshots[pulse.dataset_id]
        if mode == "context":
            first = pulse.context_start_row
            last = pulse.decay_end_row
        elif mode == "fit":
            first = pulse.fit_start_row
            last = pulse.fit_end_row
        else:
            first = pulse.decay_start_row
            last = pulse.decay_end_row
        width = 3 + len(snapshot.channels)
        if column + width - 1 > 16384 or last - first + 3 > 1048576:
            raise ValueError(
                "Selected data exceed Excel limits. Save a pulse JSON session instead."
            )
        sheet.merge_cells(
            start_row=1, start_column=column, end_row=1, end_column=column + width - 1
        )
        title = sheet.cell(1, column, pulse.label)
        title.data_type = "s"
        title.font = Font(bold=True)
        title.alignment = Alignment(horizontal="center")
        headers = ["Time (s from valve opening)", "Timestamp", "SourceRow"] + [
            f"{c.source_column} [{c.unit}]" for c in snapshot.channels
        ]
        for offset, header in enumerate(headers):
            cell = sheet.cell(2, column + offset, header)
            cell.data_type = "s"
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        keys = ["Time", "Timestamp", "SourceRow"] + [c.key for c in snapshot.channels]
        for row, values in enumerate(
            snapshot.data.iloc[first : last + 1][keys].itertuples(
                index=False, name=None
            ),
            3,
        ):
            outputs = [
                float(values[0]) - pulse.valve_time_s,
                str(values[1]),
                int(values[2]),
            ]
            outputs += [float(v) if np.isfinite(v) else None for v in values[3:]]
            for offset, value in enumerate(outputs):
                cell = sheet.cell(row, column + offset, value)
                if isinstance(value, str):
                    cell.data_type = "s"
        for offset in range(width):
            letter = sheet.cell(2, column + offset).column_letter
            sheet.column_dimensions[letter].width = 27 if offset < 2 else 24
        column += width
    sheet.freeze_panes = "A3"


def export_pulses(project, path):
    """Export the full traces, state samples, statistics and analysis setup."""
    if not project.pulses:
        raise ValueError("Capture at least one pulse first.")
    path = protect(project, path)
    if (
        sum(
            stat.total_count
            for p in project.pulses
            for state in (p.initial, p.final)
            for stat in state.values()
        )
        >= 1048575
    ):
        raise ValueError(
            "State samples exceed Excel limits. Save a pulse JSON session instead."
        )
    workbook = Workbook()
    decay = workbook.active
    decay.title = "pulse_decay"
    context = workbook.create_sheet("pulse_raw_context")
    fit = workbook.create_sheet("pulse_fit_window")
    states = workbook.create_sheet("pulse_states")
    samples = workbook.create_sheet("state_samples")
    metadata = workbook.create_sheet("selection_metadata")
    channels = workbook.create_sheet("channel_metadata")
    setup = workbook.create_sheet("analysis_setup")
    _grouped(project, decay, "decay")
    _grouped(project, context, "context")
    _grouped(project, fit, "fit")
    _append(
        states,
        [
            "Pulse ID",
            "State",
            "Data key",
            "Source column",
            "Location",
            "Quantity",
            "Canonical unit",
            "Mean",
            "Sample SD",
            "Finite n",
            "Total samples",
            "Min",
            "Max",
            "Slope per minute",
            "Max gap (s)",
            "ROI start elapsed (s)",
            "ROI end elapsed (s)",
            "Start relative to valve (s)",
            "End relative to valve (s)",
            "Start source row",
            "End source row",
            "Original unit",
            "Pressure basis",
            "Mean absolute kPa (pressure only)",
        ],
    )
    _append(
        samples,
        [
            "Pulse ID",
            "State",
            "Source column",
            "Original unit",
            "Time from valve (s)",
            "Elapsed time (s)",
            "Timestamp",
            "SourceRow",
            "Original value",
        ],
    )
    _append(
        metadata,
        [
            "Pulse ID",
            "Sample ID",
            "Source file",
            "Source SHA256",
            "Recording ID",
            "Valve opening elapsed (s)",
            "Initial pressure requested start (s)",
            "Initial pressure requested end (s)",
            "Final pressure requested start (s)",
            "Final pressure requested end (s)",
            "Initial temperature requested start (s)",
            "Initial temperature requested end (s)",
            "Final temperature requested start (s)",
            "Final temperature requested end (s)",
            "Fit start elapsed (s)",
            "Fit end elapsed (s)",
            "Fit start from valve (s)",
            "Fit end from valve (s)",
            "Context start source row",
            "Decay start source row",
            "Decay end source row",
            "Notes",
            "Review notes",
        ],
    )
    _append(
        channels,
        [
            "Recording ID",
            "Source file",
            "Data key",
            "Source column",
            "Quantity",
            "Original unit",
            "Location",
            "Import report",
        ],
    )
    _append(setup, ["Setting", "Value"])
    _append(setup, ["Pressure basis", project.pressure_basis])
    _append(setup, ["Measured barometric pressure kPa abs", project.barometric_kpa])
    _append(
        setup,
        [
            "Time reference",
            "Exact selected valve-opening time; source samples are not resampled.",
        ],
    )
    _append(
        setup,
        [
            "History matching",
            "Selected-pulse late-time and fixed-slip FDM analysis available. Save JSON sessions to retain complete analysis results.",
        ],
    )
    for key, value in project.calibrated_volumes_cm3.items():
        _append(setup, [key + " (cm3)", value])
    if project.calibration_provenance:
        _append(
            setup,
            [
                "Calibration provenance",
                json.dumps(project.calibration_provenance, ensure_ascii=False),
            ],
        )
    for snapshot in project.snapshots.values():
        for c in snapshot.channels:
            _append(
                channels,
                [
                    snapshot.dataset_id,
                    snapshot.source_path,
                    c.key,
                    c.source_column,
                    c.quantity,
                    c.unit,
                    c.location,
                    json.dumps(snapshot.import_report),
                ],
            )
    for pulse in project.pulses:
        snapshot = project.snapshots[pulse.dataset_id]
        channel_map = {c.key: c for c in snapshot.channels}
        for state, stats in [("Initial", pulse.initial), ("Final", pulse.final)]:
            for key, stat in stats.items():
                channel = channel_map[key]
                absolute = None
                if channel.quantity == "Pressure":
                    if project.pressure_basis == "Absolute":
                        absolute = stat.mean
                    elif project.barometric_kpa is not None:
                        absolute = stat.mean + project.barometric_kpa
                _append(
                    states,
                    [
                        pulse.label,
                        state,
                        key,
                        stat.source_column,
                        channel.location,
                        channel.quantity,
                        stat.unit,
                        stat.mean,
                        stat.std,
                        stat.count,
                        stat.total_count,
                        stat.minimum,
                        stat.maximum,
                        stat.slope_per_min,
                        stat.max_time_gap_s,
                        stat.start_s,
                        stat.end_s,
                        stat.start_s - pulse.valve_time_s,
                        stat.end_s - pulse.valve_time_s,
                        stat.start_source_row,
                        stat.end_source_row,
                        stat.source_unit,
                        (
                            project.pressure_basis
                            if channel.quantity == "Pressure"
                            else ""
                        ),
                        absolute,
                    ],
                )
                for t, ts, source, value in snapshot.data.iloc[
                    stat.start_row : stat.end_row + 1
                ][["Time", "Timestamp", "SourceRow", key]].itertuples(
                    index=False, name=None
                ):
                    _append(
                        samples,
                        [
                            pulse.label,
                            state,
                            channel.source_column,
                            channel.unit,
                            float(t) - pulse.valve_time_s,
                            float(t),
                            str(ts),
                            int(source),
                            float(value) if np.isfinite(value) else None,
                        ],
                    )
        _append(
            metadata,
            [
                pulse.label,
                pulse.sample_id,
                snapshot.source_path,
                snapshot.source_sha256,
                snapshot.dataset_id,
                pulse.valve_time_s,
                *pulse.initial_pressure_bounds,
                *pulse.final_pressure_bounds,
                *pulse.initial_temperature_bounds,
                *pulse.final_temperature_bounds,
                pulse.fit_start_s,
                pulse.fit_end_s,
                pulse.fit_start_s - pulse.valve_time_s,
                pulse.fit_end_s - pulse.valve_time_s,
                int(snapshot.data.SourceRow.iloc[pulse.context_start_row]),
                int(snapshot.data.SourceRow.iloc[pulse.decay_start_row]),
                int(snapshot.data.SourceRow.iloc[pulse.decay_end_row]),
                pulse.notes,
                "; ".join(pulse_review_notes(pulse, snapshot)),
            ],
        )
    for sheet in (states, samples, metadata, channels, setup):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(
                46, max(15, max(len(str(c.value or "")) for c in col[:80]) + 2)
            )
    descriptor, name = tempfile.mkstemp(dir=path.parent, suffix=".xlsx")
    os.close(descriptor)
    try:
        workbook.save(name)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)
