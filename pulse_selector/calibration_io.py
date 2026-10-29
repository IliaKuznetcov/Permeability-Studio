"""Portable calibration sessions and Excel exports. No Qt dependency."""

from __future__ import annotations
from dataclasses import asdict, fields
import json
from pathlib import Path
import tempfile
import os
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

from .calibration_models import CalibrationProject, CalibrationRun, LogSnapshot
from .models import SignalChannel
from .state_extraction import build_run, run_warnings
from .calibration import inputs_for_run, fit_project_runs, PARAMETERS

SESSION_VERSION = 1


def _protect_sources(project, path):
    path = Path(path)
    if any(
        path.resolve() == Path(s.source_path).resolve()
        for s in project.snapshots.values()
        if s.source_path
    ):
        raise ValueError(
            "Choose a new output path; source recordings cannot be overwritten."
        )
    return path


def save_session(project: CalibrationProject, path) -> None:
    path = _protect_sources(project, path)
    for run in project.runs:
        if not run.included and not run.exclusion_reason.strip():
            raise ValueError(f"{run.label}: record an exclusion reason before saving.")
    payload = {
        "format": "pulse-selector-calibration",
        "schema_version": SESSION_VERSION,
        "application_version": "4.0.1",
        "eos": project.eos,
        "snapshots": [],
        "runs": [asdict(r) for r in project.runs],
    }
    for snapshot in project.snapshots.values():
        entry = {
            f.name: getattr(snapshot, f.name)
            for f in fields(snapshot)
            if f.name not in ("data", "channels")
        }
        entry["channels"] = [asdict(c) for c in snapshot.channels]
        entry["data"] = json.loads(
            snapshot.data.to_json(orient="split", double_precision=15)
        )
        payload["snapshots"].append(entry)
    encoded = json.dumps(payload, indent=2, ensure_ascii=False, allow_nan=False)
    descriptor, name = tempfile.mkstemp(
        dir=path.parent, prefix=".calibration-", suffix=".tmp"
    )
    try:
        with os.fdopen(descriptor, "w", encoding="utf-8") as stream:
            stream.write(encoded)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)


def load_session(path) -> CalibrationProject:
    with Path(path).open(encoding="utf-8") as stream:
        payload = json.load(stream)
    if (
        payload.get("format") != "pulse-selector-calibration"
        or payload.get("schema_version") != SESSION_VERSION
    ):
        raise ValueError("Unsupported calibration-session format/version.")
    from .calibration import EOS_OPTIONS

    if payload.get("eos") not in EOS_OPTIONS:
        raise ValueError("Unsupported session equation of state.")
    project = CalibrationProject(eos=payload["eos"])
    for entry in payload["snapshots"]:
        entry = dict(entry)
        raw = entry.pop("data")
        channels = [SignalChannel(**c) for c in entry.pop("channels")]
        data = pd.DataFrame(raw["data"], columns=raw["columns"])
        required = ["Time", "Timestamp", "SourceRow"] + [c.key for c in channels]
        if data.columns.duplicated().any() or not all(c in data for c in required):
            raise ValueError("Session recording has missing or duplicate columns.")
        times = pd.to_numeric(data["Time"], errors="coerce").to_numpy(float)
        if (
            len(times) < 3
            or not np.isfinite(times).all()
            or not np.all(np.diff(times) > 0)
        ):
            raise ValueError(
                "Session recording must have increasing finite time values."
            )
        source_rows = pd.to_numeric(data["SourceRow"], errors="coerce").to_numpy(float)
        if (
            not np.isfinite(source_rows).all()
            or np.any(source_rows < 0)
            or np.any(source_rows != np.floor(source_rows))
        ):
            raise ValueError("Session source-row indices are invalid.")
        data["Time"], data["SourceRow"] = times, source_rows.astype(int)
        data["Timestamp"] = data["Timestamp"].fillna("").astype(str)
        for c in channels:
            if c.quantity not in ("Pressure", "Temperature"):
                raise ValueError("Session channel quantity is invalid.")
            data[c.key] = pd.to_numeric(data[c.key], errors="coerce").replace(
                [np.inf, -np.inf], np.nan
            )
        snapshot = LogSnapshot(channels=channels, data=data, **entry)
        if snapshot.dataset_id in project.snapshots:
            raise ValueError("Duplicate recording ID in session.")
        project.snapshots[snapshot.dataset_id] = snapshot
    for entry in payload["runs"]:
        entry = dict(entry)
        dataset_id = entry.pop("dataset_id")
        if dataset_id not in project.snapshots:
            raise ValueError("A run references a missing recording.")
        rois = entry.pop("rois")
        # Recompute statistics from the embedded samples rather than trusting cached values.
        bounds = {
            k: (v["requested_start_s"], v["requested_end_s"]) for k, v in rois.items()
        }
        channels = {k: v["channel_key"] for k, v in rois.items()}
        run = build_run(
            project.snapshots[dataset_id],
            bounds=bounds,
            channel_roles=channels,
            **entry,
        )
        if any(r.uid == run.uid or r.label == run.label for r in project.runs):
            raise ValueError("Run identifiers must be unique in a session.")
        project.runs.append(run)
    return project


def _append(sheet, values):
    # Explicit text cells also protect source names/notes beginning with '='.
    cells = []
    for value in values:
        if isinstance(value, str):
            cell = WriteOnlyCell(sheet, value=value)
            cell.data_type = "s"
            cells.append(cell)
        else:
            cells.append(value)
    sheet.append(cells)


def write_calibration_workbook(project: CalibrationProject, path) -> None:
    if not project.runs:
        raise ValueError("Capture at least one calibration run first.")
    path = _protect_sources(project, path)
    raw_count = sum(len(s.data) * len(s.channels) for s in project.snapshots.values())
    roi_count = sum(v.total_count for r in project.runs for v in r.rois.values())
    if max(raw_count, roi_count) >= 1_048_575:
        raise ValueError(
            "The trace exceeds Excel row limits. Save a JSON session to retain all samples."
        )
    workbook = Workbook()
    summary = workbook.active
    summary.title = "calibration_inputs"
    stats = workbook.create_sheet("roi_statistics")
    metadata = workbook.create_sheet("run_metadata")
    samples = workbook.create_sheet("roi_samples")
    raw_sheet = workbook.create_sheet("raw_logs")
    channels_sheet = workbook.create_sheet("source_channels")
    results = workbook.create_sheet("volume_results")
    residuals = workbook.create_sheet("fit_residuals")
    _append(
        summary,
        [
            "Include (1/0)",
            "Side",
            "Plug ID",
            "Configuration",
            "Run ID",
            "Known Vp (cm3)",
            "Pi (kPa absolute)",
            "Pb (kPa absolute)",
            "Pf (kPa absolute)",
            "Ti (degC)",
            "Tb (degC)",
            "Tf (degC)",
            "Zi used",
            "Zb used",
            "Zf used",
            "Pbar (kPa absolute)",
            "Tb source/assumption",
            "Input status",
            "ROI warnings",
            "Exclusion reason",
            "Notes",
        ],
    )
    _append(
        stats,
        ["Run ID"]
        + list(next(iter(project.runs[0].rois.values())).__dataclass_fields__),
    )
    _append(
        metadata,
        [
            "Run ID",
            "UID",
            "Recording ID",
            "Source file",
            "Source SHA256",
            "Expansion time (s)",
            "Original pressure basis",
            "Receiving-pressure mode",
            "Final temperature confirmed",
            "Pressure drift limit (kPa/min)",
            "Temperature drift limit (degC/min)",
            "EOS",
        ],
    )
    _append(
        samples,
        [
            "Run ID",
            "ROI role",
            "Source column",
            "Original unit",
            "Elapsed time (s)",
            "Timestamp",
            "SourceRow (0-based)",
            "Original value",
        ],
    )
    _append(
        raw_sheet,
        [
            "Recording ID",
            "Source file",
            "Elapsed time (s)",
            "Timestamp",
            "SourceRow (0-based)",
            "Source column",
            "Original unit",
            "Original value",
        ],
    )
    _append(
        channels_sheet,
        [
            "Recording ID",
            "Source file",
            "Time description",
            "Data key",
            "Source column",
            "Quantity",
            "Original unit",
            "Location",
            "Import report",
        ],
    )
    for snapshot in project.snapshots.values():
        for channel in snapshot.channels:
            _append(
                channels_sheet,
                [
                    snapshot.dataset_id,
                    snapshot.source_path,
                    snapshot.time_description,
                    channel.key,
                    channel.source_column,
                    channel.quantity,
                    channel.unit,
                    channel.location,
                    json.dumps(snapshot.import_report, ensure_ascii=False),
                ],
            )
            fields_ = ["Time", "Timestamp", "SourceRow", channel.key]
            for t, ts, source, value in snapshot.data[fields_].itertuples(
                index=False, name=None
            ):
                _append(
                    raw_sheet,
                    [
                        snapshot.dataset_id,
                        snapshot.source_path,
                        float(t),
                        str(ts),
                        int(source),
                        channel.source_column,
                        channel.unit,
                        float(value) if np.isfinite(value) else None,
                    ],
                )
    for run in project.runs:
        snapshot = project.snapshots[run.dataset_id]
        prefix = "Tu" if run.side == "Upstream" else "Td"
        try:
            state = inputs_for_run(run, project.eos)
            status = "Ready"
            zi, zf = state.zi, state.zf
        except ValueError as exc:
            status, zi, zf = str(exc), None, None
        bar = run.barometric_kpa
        added = bar if run.pressure_basis == "Gauge" else 0.0
        pi = run.rois["Pi"].mean + added if added is not None else None
        pf = run.rois["Pf"].mean + added if added is not None else None
        pb = bar if run.receiving_mode == "Vented" else run.receiving_pressure_kpa
        _append(
            summary,
            [
                int(run.included),
                run.side,
                run.plug_id,
                run.configuration,
                run.label,
                run.plug_volume_cm3,
                pi,
                pb,
                pf,
                run.rois[prefix + "_i"].mean,
                run.receiving_temperature_c,
                run.rois[prefix + "_f"].mean,
                zi,
                run.zb,
                zf,
                bar,
                run.receiving_temperature_source,
                status,
                "; ".join(run_warnings(run)),
                run.exclusion_reason,
                run.notes,
            ],
        )
        _append(
            metadata,
            [
                run.label,
                run.uid,
                snapshot.dataset_id,
                snapshot.source_path,
                snapshot.source_sha256,
                run.opening_time_s,
                run.pressure_basis,
                run.receiving_mode,
                int(run.final_temperature_confirmed),
                run.pressure_drift_limit_kpa_min,
                run.temperature_drift_limit_c_min,
                project.eos,
            ],
        )
        for roi in run.rois.values():
            _append(stats, [run.label] + list(asdict(roi).values()))
            for t, ts, source, value in snapshot.data.iloc[
                roi.start_row : roi.end_row + 1
            ][["Time", "Timestamp", "SourceRow", roi.channel_key]].itertuples(
                index=False, name=None
            ):
                _append(
                    samples,
                    [
                        run.label,
                        roi.role,
                        roi.source_column,
                        roi.source_unit,
                        float(t),
                        str(ts),
                        int(source),
                        float(value) if np.isfinite(value) else None,
                    ],
                )
    _append(
        results,
        [
            "Parameter",
            "Estimate (cm3)",
            "Conditional OLS SE (cm3)",
            "Approx 95% +/-1.96 SE (cm3)",
        ],
    )
    _append(
        residuals,
        [
            "Run ID",
            "x_VU",
            "x_VD",
            "x_DU",
            "x_DD",
            "RHS",
            "Fitted RHS",
            "Residual (fitted - RHS)",
            "Closure (final - initial), mmol",
        ],
    )
    try:
        fit = fit_project_runs(project.runs, project.eos)
        for name in PARAMETERS:
            se = fit.standard_errors_cm3[name]
            _append(
                results,
                [
                    name,
                    fit.volumes_cm3[name],
                    se,
                    se * 1.96 if se is not None else None,
                ],
            )
        for name, value in [
            ("VU+DU", fit.volumes_cm3["VU"] + fit.volumes_cm3["DU"]),
            ("VD+DD", fit.volumes_cm3["VD"] + fit.volumes_cm3["DD"]),
            ("DU+DD", fit.volumes_cm3["DU"] + fit.volumes_cm3["DD"]),
            ("Included runs", len(fit.labels)),
            ("SSE", fit.sse),
            ("RMSE", fit.rmse),
            ("R squared (included rows)", fit.r_squared),
            ("Matrix rank", fit.rank),
            ("Condition number", fit.condition_number),
            ("Residual degrees of freedom", fit.degrees_of_freedom),
            ("EOS", fit.eos),
        ]:
            _append(results, [name, value])
        for message in fit.warnings:
            _append(results, ["Review note", message])
        for i, label in enumerate(fit.labels):
            _append(
                residuals,
                [label]
                + fit.matrix[i]
                + [
                    fit.rhs[i],
                    fit.fitted_rhs[i],
                    fit.residuals[i],
                    fit.closure_mmol[i],
                ],
            )
    except ValueError as exc:
        _append(results, ["Fit not available", str(exc)])
    _append(
        results,
        [
            "Method",
            "Unweighted least squares by SVD; SE is conditional on X, not full measurement uncertainty.",
        ],
    )
    _append(
        results,
        [
            "Pressure and temperature",
            "All gas-law pressures are absolute kPa; temperatures are kelvin internally.",
        ],
    )
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="254C64")
        for column in sheet.columns:
            longest = max(
                len(str(c.value if c.value is not None else "")) for c in column[:100]
            )
            sheet.column_dimensions[column[0].column_letter].width = min(
                max(12, longest + 2), 48
            )
    # Atomic replacement prevents leaving a half-written export on an interrupted write.
    descriptor, name = tempfile.mkstemp(dir=path.parent, suffix=".xlsx")
    os.close(descriptor)
    try:
        workbook.save(name)
        os.replace(name, path)
    finally:
        Path(name).unlink(missing_ok=True)
