"""Excel serialization of the model; independent of Qt."""

from pathlib import Path
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from .models import ApplicationModel


def write_workbook(model: ApplicationModel, path: Path) -> None:
    """Export a variable-width group per pulse, using original sample values."""
    if model.data is None or not model.selections:
        raise ValueError("Capture at least one pulse before exporting.")
    path = Path(path)
    if model.source_path and path.resolve() == model.source_path.resolve():
        raise ValueError(
            "Choose a new export filename to preserve the source recording."
        )
    workbook = Workbook()
    pulse_sheet = workbook.active
    pulse_sheet.title = "pulse_decay"
    context_sheet = workbook.create_sheet("pulse_raw_context")
    metadata_sheet = workbook.create_sheet("selection_metadata")
    channels_sheet = workbook.create_sheet("channel_metadata")
    plateau_sheet = workbook.create_sheet("plateau_summary")
    plateau_samples_sheet = workbook.create_sheet("pulse_plateau")
    for sheet, context in [(pulse_sheet, False), (context_sheet, True)]:
        _write_grouped_pulses(model, sheet, context)
    _write_grouped_pulses(model, plateau_samples_sheet, plateau=True)
    metadata_headers = [
        "Pulse",
        "Source file",
        "Date column",
        "Time column",
        "Decay start elapsed time (s)",
        "Decay end elapsed time (s)",
        "Decay duration (s)",
        "Decay start timestamp",
        "Decay end timestamp",
        "Decay start cleaned row",
        "Decay end cleaned row",
        "Raw-context start elapsed time (s)",
        "Raw-context start timestamp",
        "Raw-context start cleaned row",
        "Start source row (0-based)",
        "End source row (0-based)",
        "Context source row (0-based)",
        "Isolated date rows repaired",
        "Invalid-time rows removed",
        "Non-increasing rows removed",
        "Plateau start elapsed time (s)",
        "Plateau end elapsed time (s)",
        "Plateau duration (s)",
        "Plateau start timestamp",
        "Plateau end timestamp",
        "Plateau start cleaned row",
        "Plateau end cleaned row",
        "Plateau start source row (0-based)",
        "Plateau end source row (0-based)",
    ]
    metadata_sheet.append(metadata_headers)
    for selection in model.selections:
        metadata_sheet.append(
            [
                selection.number,
                str(model.source_path or ""),
                model.date_column or "",
                model.time_column,
                selection.start_time_s,
                selection.end_time_s,
                selection.end_time_s - selection.start_time_s,
                selection.start_timestamp,
                selection.end_timestamp,
                selection.start_row,
                selection.end_row,
                selection.context_start_time_s,
                selection.context_start_timestamp,
                selection.context_start_row,
                int(model.data["SourceRow"].iloc[selection.start_row]),
                int(model.data["SourceRow"].iloc[selection.end_row]),
                int(model.data["SourceRow"].iloc[selection.context_start_row]),
                model.import_report["repaired_date_rows"],
                model.import_report["invalid_time_rows"],
                model.import_report["nonincreasing_rows"],
                selection.plateau.start_time_s,
                selection.plateau.end_time_s,
                selection.plateau.end_time_s - selection.plateau.start_time_s,
                selection.plateau.start_timestamp,
                selection.plateau.end_timestamp,
                selection.plateau.start_row,
                selection.plateau.end_row,
                int(model.data["SourceRow"].iloc[selection.plateau.start_row]),
                int(model.data["SourceRow"].iloc[selection.plateau.end_row]),
            ]
        )
    channels_sheet.append(
        [
            "Data key",
            "Source column",
            "Quantity",
            "Original unit",
            "Location",
            "Missing samples in log",
        ]
    )
    for channel in model.channels:
        channels_sheet.append(
            [
                channel.key,
                channel.source_column,
                channel.quantity,
                channel.unit,
                channel.location,
                model.import_report["missing_samples"][channel.key],
            ]
        )
    plateau_sheet.append(
        [
            "Pulse",
            "Data key",
            "Source column",
            "Quantity",
            "Original unit",
            "Location",
            "Plateau mean",
            "Plateau sample SD (ddof=1)",
            "Valid samples",
            "Total interval samples",
            "Missing samples",
            "Start elapsed time (s)",
            "End elapsed time (s)",
            "Start timestamp",
            "End timestamp",
            "Method",
        ]
    )
    for selection in model.selections:
        plateau = selection.plateau
        total = plateau.end_row - plateau.start_row + 1
        for channel in model.channels:
            stat = plateau.values[channel.key]
            plateau_sheet.append(
                [
                    selection.number,
                    channel.key,
                    channel.source_column,
                    channel.quantity,
                    channel.unit,
                    channel.location,
                    stat.mean,
                    stat.std,
                    stat.count,
                    total,
                    total - stat.count,
                    plateau.start_time_s,
                    plateau.end_time_s,
                    plateau.start_timestamp,
                    plateau.end_timestamp,
                    "Arithmetic mean of finite samples in manually selected inclusive interval",
                ]
            )
    for sheet in (metadata_sheet, channels_sheet, plateau_sheet):
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAF7")
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = max(len(str(cell.value or "")) for cell in column) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(width, 42)
    workbook.save(path)


def _write_grouped_pulses(
    model, worksheet, include_context: bool = False, plateau: bool = False
) -> None:
    headers = ["Time", "Timestamp", "SourceRow"] + [
        c.source_column for c in model.channels
    ]
    width = len(headers)
    for selection in model.selections:
        start_row = (
            selection.plateau.start_row
            if plateau
            else selection.context_start_row if include_context else selection.start_row
        )
        end_row = selection.plateau.end_row if plateau else selection.end_row
        frame = model.data.iloc[start_row : end_row + 1]
        start_column = 1 + width * (selection.number - 1)
        worksheet.merge_cells(
            start_row=1,
            start_column=start_column,
            end_row=1,
            end_column=start_column + width - 1,
        )
        title = worksheet.cell(1, start_column, f"Pulse {selection.number}")
        title.font = Font(bold=True)
        title.alignment = Alignment(horizontal="center")
        title.fill = PatternFill("solid", fgColor="D9EAD3")
        for offset, header in enumerate(headers):
            cell = worksheet.cell(2, start_column + offset, header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EAF4E6")
        fields = ["Time", "Timestamp", "SourceRow"] + [c.key for c in model.channels]
        for output_row, values in enumerate(
            frame[fields].itertuples(index=False, name=None), start=3
        ):
            output = [
                float(values[0]) - selection.start_time_s,
                str(values[1]),
                int(values[2]),
            ]
            output += [float(v) if np.isfinite(v) else None for v in values[3:]]
            for offset, value in enumerate(output):
                worksheet.cell(output_row, start_column + offset, value)
        for offset, header in enumerate(headers):
            letter = worksheet.cell(2, start_column + offset).column_letter
            worksheet.column_dimensions[letter].width = (
                23 if header == "Timestamp" else max(16, len(header) + 2)
            )
    worksheet.freeze_panes = "A3"
